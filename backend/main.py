"""
FastAPI application — Building Bill Manager
Endpoints:
  POST /api/auth/login           — get JWT token
  GET  /api/buildings            — building + room config
  GET  /api/months               — list all months
  POST /api/months               — create a new month (seeds carry-forward)
  GET  /api/months/{year}/{month}
  PUT  /api/months/{year}/{month}/rate    — update electricity rate settings
  GET  /api/months/{year}/{month}/rooms
  PUT  /api/months/{year}/{month}/rooms  — bulk-save room readings
  POST /api/months/{year}/{month}/generate-pdf
"""
import os
import datetime
import calendar
from typing import Annotated

from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, FileResponse
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.staticfiles import StaticFiles

from jose import JWTError, jwt
import bcrypt as _bcrypt
from pydantic import BaseModel
from sqlalchemy.orm import Session

from .database import engine, get_db, Base
from . import models
from .config import (
    APP_PASSWORD,
    SECRET_KEY,
    TOKEN_EXPIRE_HOURS,
    BUILDINGS,
    DEFAULT_ELECTRICITY_RATE,
    DEFAULT_ELECTRICITY_RATE_HIGH,
    DEFAULT_RATE_THRESHOLD,
    get_room_defaults,
)
from .pdf_gen import generate_pdf_for_month

# ── Create tables ──────────────────────────────────────────────────────────────
Base.metadata.create_all(bind=engine)

app = FastAPI(title="Building Bill Manager", version="1.0.0")

# ── CORS ───────────────────────────────────────────────────────────────────────
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Auth ───────────────────────────────────────────────────────────────────────
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login")
ALGORITHM = "HS256"

HASHED_PASSWORD = _bcrypt.hashpw(APP_PASSWORD.encode(), _bcrypt.gensalt())


def create_token(data: dict) -> str:
    to_encode = data.copy()
    expire = datetime.datetime.utcnow() + datetime.timedelta(hours=TOKEN_EXPIRE_HOURS)
    to_encode["exp"] = expire
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)


def get_current_user(token: Annotated[str, Depends(oauth2_scheme)]):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED)
        return username
    except JWTError:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")


# ── Pydantic schemas ───────────────────────────────────────────────────────────
class Token(BaseModel):
    access_token: str
    token_type: str


class RoomReadingIn(BaseModel):
    building_num: int
    room_no: str
    meter_no: str = ""
    previous_units: float = 0
    present_units: float = 0
    rent: float = 0
    gas_bill: float = 1080
    service_charge: float = 500
    previous_due: float = 0
    paid: float = 0
    comments: str = ""


class RateSettingsIn(BaseModel):
    electricity_rate: float
    electricity_rate_high: float
    rate_threshold_units: int


class MonthCreateIn(BaseModel):
    year: int
    month: int  # 1–12


# ── Helper: compute derived fields for a room reading ─────────────────────────
def _room_to_dict(r: models.RoomReading, rec: models.MonthRecord) -> dict:
    units = r.units_used
    eb = r.electric_bill(rec.electricity_rate, rec.electricity_rate_high, rec.rate_threshold_units)
    tb = r.total_bill(rec.electricity_rate, rec.electricity_rate_high, rec.rate_threshold_units)
    return {
        "id": r.id,
        "building_num": r.building_num,
        "room_no": r.room_no,
        "meter_no": r.meter_no,
        "previous_units": r.previous_units,
        "present_units": r.present_units,
        "units_used": units,
        "electric_bill": round(eb, 2),
        "gas_bill": r.gas_bill,
        "service_charge": r.service_charge,
        "rent": r.rent,
        "previous_due": r.previous_due,
        "paid": r.paid,
        "comments": r.comments,
        "total_bill": round(tb, 2),
    }


# ══════════════════════════════════════════════════════════════════════════════
# Auth routes
# ══════════════════════════════════════════════════════════════════════════════
@app.post("/api/auth/login", response_model=Token)
def login(form_data: Annotated[OAuth2PasswordRequestForm, Depends()]):
    if not _bcrypt.checkpw(form_data.password.encode(), HASHED_PASSWORD):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Wrong password")
    token = create_token({"sub": "admin"})
    return {"access_token": token, "token_type": "bearer"}


# ══════════════════════════════════════════════════════════════════════════════
# Building config route
# ══════════════════════════════════════════════════════════════════════════════
@app.get("/api/buildings")
def get_buildings(user: str = Depends(get_current_user)):
    return BUILDINGS


# ══════════════════════════════════════════════════════════════════════════════
# Month routes
# ══════════════════════════════════════════════════════════════════════════════
@app.get("/api/months")
def list_months(db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    records = db.query(models.MonthRecord).order_by(
        models.MonthRecord.year.desc(), models.MonthRecord.month.desc()
    ).all()
    return [
        {
            "year": r.year,
            "month": r.month,
            "month_name": calendar.month_name[r.month],
            "electricity_rate": r.electricity_rate,
            "electricity_rate_high": r.electricity_rate_high,
            "rate_threshold_units": r.rate_threshold_units,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in records
    ]


@app.post("/api/months", status_code=201)
def create_month(body: MonthCreateIn, db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    existing = db.query(models.MonthRecord).filter_by(year=body.year, month=body.month).first()
    if existing:
        raise HTTPException(400, f"{calendar.month_name[body.month]} {body.year} already exists.")

    # Find previous month's record for carry-forward
    prev_date = datetime.date(body.year, body.month, 1) - datetime.timedelta(days=1)
    prev_record = db.query(models.MonthRecord).filter_by(
        year=prev_date.year, month=prev_date.month
    ).first()

    new_record = models.MonthRecord(
        year=body.year,
        month=body.month,
        electricity_rate=DEFAULT_ELECTRICITY_RATE,
        electricity_rate_high=DEFAULT_ELECTRICITY_RATE_HIGH,
        rate_threshold_units=DEFAULT_RATE_THRESHOLD,
    )
    db.add(new_record)
    db.flush()  # get new_record.id

    # Seed rooms from previous month (carry-forward) or config defaults
    prev_readings: dict[tuple, dict] = {}
    if prev_record:
        for pr in prev_record.rooms:
            prev_readings[(pr.building_num, pr.room_no)] = {
                "present_units": pr.present_units,
                "rent": pr.rent,
                "gas_bill": pr.gas_bill,
                "service_charge": pr.service_charge,
            }

    for bldg in BUILDINGS:
        for room in bldg["rooms"]:
            defaults = get_room_defaults(bldg["id"], room["room_no"])
            key = (bldg["id"], room["room_no"])
            prev = prev_readings.get(key)
            prev_u = prev["present_units"] if prev else 0.0
            rr = models.RoomReading(
                month_record_id=new_record.id,
                building_num=bldg["id"],
                room_no=room["room_no"],
                meter_no=defaults.get("meter_no", ""),
                previous_units=prev_u,
                present_units=prev_u,  # start equal; user updates present
                rent=prev["rent"] if prev else defaults.get("rent", 0),
                gas_bill=prev["gas_bill"] if prev else defaults.get("gas_bill", 1080),
                service_charge=prev["service_charge"] if prev else defaults.get("service_charge", 500),
            )
            db.add(rr)

    db.commit()
    db.refresh(new_record)
    return {"year": new_record.year, "month": new_record.month, "id": new_record.id}


@app.get("/api/months/{year}/{month}")
def get_month(year: int, month: int, db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")
    return {
        "year": rec.year,
        "month": rec.month,
        "month_name": calendar.month_name[rec.month],
        "electricity_rate": rec.electricity_rate,
        "electricity_rate_high": rec.electricity_rate_high,
        "rate_threshold_units": rec.rate_threshold_units,
        "rooms": [_room_to_dict(r, rec) for r in rec.rooms],
    }


@app.put("/api/months/{year}/{month}/rate")
def update_rate(
    year: int, month: int, body: RateSettingsIn,
    db: Session = Depends(get_db), user: str = Depends(get_current_user)
):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")
    rec.electricity_rate = body.electricity_rate
    rec.electricity_rate_high = body.electricity_rate_high
    rec.rate_threshold_units = body.rate_threshold_units
    rec.updated_at = datetime.datetime.utcnow()
    db.commit()
    return {"ok": True}


@app.get("/api/months/{year}/{month}/rooms")
def get_rooms(year: int, month: int, db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")
    return [_room_to_dict(r, rec) for r in rec.rooms]


@app.put("/api/months/{year}/{month}/rooms")
def update_rooms(
    year: int,
    month: int,
    rooms: list[RoomReadingIn],
    db: Session = Depends(get_db),
    user: str = Depends(get_current_user),
):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")

    for incoming in rooms:
        rr = (
            db.query(models.RoomReading)
            .filter_by(month_record_id=rec.id, building_num=incoming.building_num, room_no=incoming.room_no)
            .first()
        )
        if not rr:
            rr = models.RoomReading(month_record_id=rec.id, building_num=incoming.building_num, room_no=incoming.room_no)
            db.add(rr)

        rr.meter_no = incoming.meter_no
        rr.previous_units = incoming.previous_units
        rr.present_units = incoming.present_units
        rr.rent = incoming.rent
        rr.gas_bill = incoming.gas_bill
        rr.service_charge = incoming.service_charge
        rr.previous_due = incoming.previous_due
        rr.paid = incoming.paid
        rr.comments = incoming.comments

    rec.updated_at = datetime.datetime.utcnow()
    db.commit()
    db.refresh(rec)
    return [_room_to_dict(r, rec) for r in rec.rooms]


@app.post("/api/months/{year}/{month}/generate-pdf")
def generate_pdf(year: int, month: int, two_up: bool = False, db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")

    rooms_data = [_room_to_dict(r, rec) for r in rec.rooms]
    pdf_bytes = generate_pdf_for_month(rooms_data, year, month, two_up)

    month_name = calendar.month_name[month]
    filename = f"bills_{month_name}_{year}{'_2up' if two_up else ''}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.delete("/api/months/{year}/{month}")
def delete_month(year: int, month: int, db: Session = Depends(get_db), user: str = Depends(get_current_user)):
    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")
    db.delete(rec)  # cascade deletes room readings
    db.commit()
    return {"ok": True, "message": f"{calendar.month_name[month]} {year} deleted"}


# ══════════════════════════════════════════════════════════════════════════════
# CSV Import
# ══════════════════════════════════════════════════════════════════════════════
import csv
import io
from fastapi import UploadFile, File

@app.post("/api/months/import-csv")
async def import_csv(
    year: int,
    month: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: str = Depends(get_current_user),
):
    """
    Import a CSV file in the existing format.
    The CSV has two building sections separated by blank rows:
      - First section:  House 220 rooms  (rows until first blank)
      - Second section: House 226 rooms  (rows after blank rows)
    
    Column mapping (0-indexed):
      0: Room No.
      1: Meter No.
      2: [previous] Units T.
      3: present units
      4: [present] UNITS  (units used, computed — we ignore and recompute)
      5: [present] BILL   (electric bill, computed — we ignore)
      6: Rent
      7: Gas Bill
      8: Service Charge
      9: Total W BILLS (computed — ignore)
      10: Previous Due
      11: Paid
      12: Due (computed — ignore)
      13: Occupied (flag — ignore)
      14: TOTAL W. OC (computed — ignore)
      15: advance paid
    """
    # Check if month already exists
    existing = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if existing:
        raise HTTPException(400, f"{calendar.month_name[month]} {year} already exists. Delete it first or choose a different month.")

    content = await file.read()
    text = content.decode("utf-8-sig")  # handle BOM from Excel
    reader = csv.reader(io.StringIO(text))

    rows = list(reader)
    if not rows:
        raise HTTPException(400, "CSV file is empty")

    # Skip header row
    data_rows = rows[1:]

    # Split into building sections by blank rows
    building_configs = {b["id"]: b for b in BUILDINGS}
    building_ids = [b["id"] for b in BUILDINGS]  # [220, 226]

    sections: list[list[list[str]]] = []
    current_section: list[list[str]] = []

    for row in data_rows:
        # A row is "blank" if Room No. (col 0) is empty
        room_no_val = row[0].strip() if len(row) > 0 else ""
        if not room_no_val:
            if current_section:
                sections.append(current_section)
                current_section = []
            continue
        current_section.append(row)

    if current_section:
        sections.append(current_section)

    if len(sections) == 0:
        raise HTTPException(400, "No room data found in the CSV")

    # Create the month record
    new_record = models.MonthRecord(
        year=year,
        month=month,
        electricity_rate=DEFAULT_ELECTRICITY_RATE,
        electricity_rate_high=DEFAULT_ELECTRICITY_RATE_HIGH,
        rate_threshold_units=DEFAULT_RATE_THRESHOLD,
    )
    db.add(new_record)
    db.flush()

    imported_count = 0

    def safe_float(val: str, default: float = 0.0) -> float:
        try:
            return float(val.strip().replace(",", ""))
        except (ValueError, AttributeError):
            return default

    for section_idx, section in enumerate(sections):
        # Map section index to building ID
        if section_idx < len(building_ids):
            bldg_id = building_ids[section_idx]
        else:
            break  # only handle known buildings

        bldg_config = building_configs[bldg_id]

        for row in section:
            if len(row) < 7:
                continue

            room_no = row[0].strip()
            if not room_no:
                continue

            # Look up defaults from config
            defaults = get_room_defaults(bldg_id, room_no)
            meter_no = row[1].strip() if row[1].strip() else defaults.get("meter_no", "")

            rr = models.RoomReading(
                month_record_id=new_record.id,
                building_num=bldg_id,
                room_no=room_no,
                meter_no=meter_no,
                previous_units=safe_float(row[2]) if len(row) > 2 else 0,
                present_units=safe_float(row[3]) if len(row) > 3 else 0,
                rent=safe_float(row[6]) if len(row) > 6 else defaults.get("rent", 0),
                gas_bill=safe_float(row[7]) if len(row) > 7 else bldg_config["default_gas_bill"],
                service_charge=safe_float(row[8]) if len(row) > 8 else bldg_config["default_service_charge"],
                previous_due=safe_float(row[10]) if len(row) > 10 else 0,
                paid=safe_float(row[15]) if len(row) > 15 else 0,
            )
            db.add(rr)
            imported_count += 1

    db.commit()
    db.refresh(new_record)

    return {
        "message": f"Imported {imported_count} rooms into {calendar.month_name[month]} {year}",
        "year": new_record.year,
        "month": new_record.month,
        "rooms_imported": imported_count,
    }


# ══════════════════════════════════════════════════════════════════════════════
# CSV / Excel Export
# ══════════════════════════════════════════════════════════════════════════════
@app.get("/api/months/{year}/{month}/export")
def export_month(
    year: int,
    month: int,
    format: str = "csv",
    db: Session = Depends(get_db),
    user: str = Depends(get_current_user),
):
    """
    Export a month's room data as CSV or Excel.
    CSV format matches the import format exactly — exported files can be
    re-imported via the Dashboard's "Import CSV" button.
    """
    if format not in ("csv", "xlsx"):
        raise HTTPException(400, "Format must be 'csv' or 'xlsx'")

    rec = db.query(models.MonthRecord).filter_by(year=year, month=month).first()
    if not rec:
        raise HTTPException(404, "Month not found")

    rooms = rec.rooms

    # Group by building
    building_220 = [r for r in rooms if r.building_num == 220]
    building_226 = [r for r in rooms if r.building_num == 226]

    # Sort rooms within each building by room_no
    building_220.sort(key=lambda r: r.room_no)
    building_226.sort(key=lambda r: r.room_no)

    month_name = calendar.month_name[month]

    def room_export_row(r):
        units_used = r.units_used
        eb = r.electric_bill(rec.electricity_rate, rec.electricity_rate_high, rec.rate_threshold_units)
        total_w_bills = eb + r.gas_bill + r.service_charge + r.rent
        total_bill = total_w_bills + r.previous_due - r.paid
        return [
            r.room_no,                     # 0: Room No.
            r.meter_no,                    # 1: Meter No.
            r.previous_units,              # 2: Previous Units
            r.present_units,               # 3: Present Units
            units_used,                    # 4: Units Used
            round(eb, 2),                  # 5: Electric Bill
            r.rent,                        # 6: Rent
            r.gas_bill,                    # 7: Gas Bill
            r.service_charge,              # 8: Service Charge
            round(total_w_bills, 2),       # 9: Total W Bills
            r.previous_due,                # 10: Previous Due
            0,                             # 11: Paid (import reads from col 15)
            round(total_bill, 2),          # 12: Due
            "",                            # 13: Occupied
            round(total_bill, 2),          # 14: TOTAL W. OC
            r.paid,                        # 15: Advance Paid (import reads paid from here)
        ]

    header = [
        "Room No.", "Meter No.", "Previous Units", "Present Units",
        "Units Used", "Electric Bill", "Rent", "Gas Bill",
        "Service Charge", "Total W Bills", "Previous Due", "Paid",
        "Due", "Occupied", "TOTAL W. OC", "Advance Paid",
    ]

    filename = f"bills_{month_name}_{year}"

    if format == "xlsx":
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{month_name} {year}"

        # Header row
        for col_idx, h in enumerate(header, 1):
            ws.cell(row=1, column=col_idx, value=h)

        row_num = 2

        # Building 220 rooms
        for r in building_220:
            row_data = room_export_row(r)
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_idx, value=val)
            row_num += 1

        # Blank separator row
        row_num += 1

        # Building 226 rooms
        for r in building_226:
            row_data = room_export_row(r)
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_idx, value=val)
            row_num += 1

        # Auto-fit column widths
        for col_idx in range(1, len(header) + 1):
            max_width = len(str(header[col_idx - 1]))
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2, max_row=row_num - 1):
                for cell in row:
                    if cell.value is not None:
                        max_width = max(max_width, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 2, 30)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
        )
    else:
        # CSV format
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(header)

        for r in building_220:
            writer.writerow(room_export_row(r))

        writer.writerow([])  # blank separator row between buildings

        for r in building_226:
            writer.writerow(room_export_row(r))

        output.seek(0)

        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
        )


# ── Serve frontend build (production) ─────────────────────────────────────────
# ── Serve frontend build (production) ─────────────────────────────────────────
FRONTEND_DIST = os.path.join(os.path.dirname(__file__), "..", "frontend", "dist")

if os.path.isdir(FRONTEND_DIST):
    # Serve static assets (_assets, vite.svg, etc)
    app.mount("/assets", StaticFiles(directory=os.path.join(FRONTEND_DIST, "assets")), name="assets")
    
    # Catch-all route to serve index.html for React Router SPA
    @app.get("/{full_path:path}")
    async def serve_frontend(full_path: str):
        # Prevent shadowing the API routes
        if full_path.startswith("api/"):
            raise HTTPException(status_code=404, detail="API route not found")
            
        # Serve actual files if requested (e.g. vite.svg)
        file_path = os.path.join(FRONTEND_DIST, full_path)
        if os.path.isfile(file_path):
            return FileResponse(file_path)
            
        # Otherwise fallback to index.html for React Router
        return FileResponse(os.path.join(FRONTEND_DIST, "index.html"))
